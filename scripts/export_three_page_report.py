from pathlib import Path
from html import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parent.parent
DOCX_PATH = ROOT / 'PROJECT_CODE_3_PAGES.docx'
MD_PATH = ROOT / 'PROJECT_CODE_3_PAGES.md'
HTML_PATH = ROOT / 'PROJECT_CODE_3_PAGES.html'
PDF_PATH = ROOT / 'PROJECT_CODE_3_PAGES.pdf'
PAGES = [
    {
        'title': '1. Image Analysis and Prediction',
        'summary': 'The app runs the optional classifier and independently reads color bands. When they disagree, a valid color-band value is displayed and the result is flagged for review.',
        'blocks': [
            ('src/App.tsx - classifier and color-band cross-check', """const modelResult = await detectWithModel(image)

if (modelResult?.type === 'classification') {
  let bandCheck = null
  try {
    bandCheck = detectResistorBandsFromImage(image)
  } catch {
    bandCheck = null
  }

  const confidence = Math.round(modelResult.confidence * 100)
  const predictedOhms = parseResistanceLabel(modelResult.label)
  const disagrees = bandCheck && predictedOhms !== null &&
    Math.abs(predictedOhms - bandCheck.result.ohms) >
      Math.max(0.01, predictedOhms * 0.01)
  const bandEstimate = bandCheck?.result.formatted ?? ''

  setImageState({
    status: disagrees || !bandCheck ? 'warning' : 'success',
    bands: bandCheck?.bands ?? [],
    resistance: bandEstimate || modelResult.label,
    tolerance: bandCheck ? `+/-${bandCheck.result.tolerance}%` : '',
    bandEstimate,
    confidence,
    imageUrl: dataUrl,
  })
  return
}"""),
        ],
    },
    {
        'title': '2. Color-Band Detection and Value Calculation',
        'summary': 'The vision helper samples vertical colors inside the resistor body, groups stripe runs, and reads the sequence in either direction. The decoder rejects gold and silver in digit positions.',
        'blocks': [
            ('src/bandVision.ts - RGB hue classification', """const saturation = maximum === 0 ? 0 : (maximum - minimum) / maximum
if (saturation < 0.42) {
  if (value < 0.18) return 'black'
  if (saturation < 0.1 && value > 0.88) return 'white'
  if (saturation < 0.16 && value > 0.55) return 'silver'
  return null
}

const delta = (maximum - minimum) / 255
let hue = 0
if (maximum === red) {
  hue = 60 * (((greenNorm - blueNorm) / delta) % 6)
} else if (maximum === green) {
  hue = 60 * ((blueNorm - redNorm) / delta + 2)
} else {
  hue = 60 * ((redNorm - greenNorm) / delta + 4)
}
hue = (hue + 360) % 360

if (hue < 10 || hue >= 350) return 'red'
if (hue < 35) return value < 0.62 ? 'brown' : 'orange'
if (hue < 49) return 'gold'
if (hue < 75) return 'yellow'
if (hue < 165) return 'green'
if (hue < 235) return 'blue'
return 'violet'"""),
            ('src/resistorCalculator.ts - decode four and five bands', """export function decodeResistor(bands: string[], type: BandType) {
  const normalized = bands.map((band) => band.toLowerCase() as ResistorColor)
  const expectedCount = type === '4-band' ? 4 : 5
  if (normalized.length !== expectedCount ||
      normalized.some((color) => !(color in colorValues))) {
    throw new Error(`A ${type} resistor requires ${expectedCount} valid color bands.`)
  }

  if (type === '4-band') {
    const [a, b, multiplier, toleranceBand] = normalized
    if (!digitColors.has(a) || !digitColors.has(b)) {
      throw new Error('Gold and silver cannot be significant-digit bands.')
    }
    const ohms = (colorValues[a] * 10 + colorValues[b]) * colorMultiplier[multiplier]
    const tolerance = toleranceBand === 'gold' ? 5 : toleranceBand === 'silver' ? 10 : 20
    return { ohms, tolerance, formatted: formatResistance(ohms, tolerance) }
  }

  const [a, b, c, multiplier, toleranceBand] = normalized
  if (!digitColors.has(a) || !digitColors.has(b) || !digitColors.has(c)) {
    throw new Error('Gold and silver cannot be significant-digit bands.')
  }
  const ohms = (colorValues[a] * 100 + colorValues[b] * 10 + colorValues[c]) *
    colorMultiplier[multiplier]
  const tolerance = toleranceBand === 'gold' ? 5 : toleranceBand === 'silver' ? 10 : 20
  return { ohms, tolerance, formatted: formatResistance(ohms, tolerance) }
}"""),
        ],
    },
    {
        'title': '3. Training, Inference, and Test Results',
        'summary': 'The YOLO11 classifier is trained from the labeled resistor images. At runtime, Python returns the class label and confidence; pixel-based band decoding provides an independent value check.',
        'blocks': [
            ('scripts/train_resistor_classifier.py - model training and evaluation', """model = YOLO('yolo11n-cls.pt')
model.train(
    data=str(data_dir),
    epochs=args.epochs,
    imgsz=args.imgsz,
    batch=args.batch,
    patience=10,
    seed=args.seed,
    project=str(PROJECT_ROOT / 'runs' / 'classify'),
    name='resistor-values',
    exist_ok=True,
)

best_weights = PROJECT_ROOT / 'runs' / 'classify' / 'resistor-values' / 'weights' / 'best.pt'
weights_path.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(best_weights, weights_path)
metrics = YOLO(str(weights_path)).val(data=str(data_dir), split='test')
print(f'Test top-1 accuracy: {metrics.top1:.4f}')"""),
            ('scripts/detect_resistor.py - runtime model prediction', """model = YOLO(resolved_model)
image_size = 224 if model.task == 'classify' else 640
results = model(str(image_path), imgsz=image_size, verbose=False)

for result in results:
    if result.probs is not None:
        class_id = int(result.probs.top1)
        classification = {
            'label': str(result.names[class_id]),
            'confidence': float(result.probs.top1conf),
        }

print(json.dumps({
    'detections': detections,
    'classification': classification,
}))"""),
            ('src/bandVision.test.ts - the two resistor patterns', """expect(detectColorBandSequence(makeResistorImage())).toEqual(
  ['yellow', 'violet', 'red', 'gold'],
)
expect(detectColorBandSequence(makeResistorImage([
  [210, 25, 24], [210, 25, 24], [124, 63, 15], [204, 160, 35],
]))).toEqual(['red', 'red', 'brown', 'gold'])"""),
        ],
        'note': 'Classifier test result: 91% top-1 accuracy and 99.7% top-5 accuracy. The two sample color codes correspond to 4.7 kOhm and 220 ohms.',
    },
]


def add_code(document: Document, heading: str, code: str) -> None:
    label = document.add_paragraph()
    label.paragraph_format.space_before = Pt(4)
    label.paragraph_format.space_after = Pt(2)
    label_run = label.add_run(heading)
    label_run.bold = True
    label_run.font.size = Pt(8)
    label_run.font.color.rgb = RGBColor(31, 78, 121)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.1)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = Pt(8)
    run = paragraph.add_run(code.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(6.8)


def write_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)

    normal = document.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(8)
    normal.paragraph_format.space_after = Pt(3)

    for index, page in enumerate(PAGES):
        if index:
            document.add_page_break()

        heading = document.add_paragraph()
        heading.paragraph_format.space_after = Pt(3)
        run = heading.add_run(page['title'])
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(31, 78, 121)

        if index == 0:
            subtitle = document.add_paragraph('RESISTOR VALUE IDENTIFIER | CONCISE CODE REPORT')
            subtitle.paragraph_format.space_after = Pt(4)
            for run in subtitle.runs:
                run.bold = True
                run.font.size = Pt(7.5)

        document.add_paragraph(page['summary'])
        for block_heading, code in page['blocks']:
            add_code(document, block_heading, code)
        if page.get('note'):
            note = document.add_paragraph(page['note'])
            note.paragraph_format.space_before = Pt(4)
            for run in note.runs:
                run.bold = True
                run.font.size = Pt(8)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Resistor Value Identifier | Three-page code summary').font.size = Pt(7)
    document.save(DOCX_PATH)


def write_markdown() -> None:
    sections = ['# Resistor Value Identifier - Three-Page Code Summary', '']
    for index, page in enumerate(PAGES):
        if index:
            sections.extend(['', '<!-- PAGE BREAK -->', ''])
        sections.extend([f"## {page['title']}", '', page['summary'], ''])
        for heading, code in page['blocks']:
            sections.extend([f'### {heading}', '', '```text', code.strip(), '```', ''])
        if page.get('note'):
            sections.extend([page['note'], ''])
    MD_PATH.write_text('\n'.join(sections), encoding='utf-8')


def write_html() -> None:
    page_html = []
    for index, page in enumerate(PAGES):
        blocks = ''.join(
            f'<h2>{escape(heading)}</h2><pre><code>{escape(code.strip())}</code></pre>'
            for heading, code in page['blocks']
        )
        note = f'<p class="note">{escape(page["note"])}</p>' if page.get('note') else ''
        cover = '<p class="cover">RESISTOR VALUE IDENTIFIER | CONCISE CODE REPORT</p>' if index == 0 else ''
        page_html.append(
            f'<article class="page">{cover}<h1>{escape(page["title"])}</h1>'
            f'<p class="summary">{escape(page["summary"])}</p>{blocks}{note}'
            f'<footer>Resistor Value Identifier | Page {index + 1} of 3</footer></article>'
        )

    html = '''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Resistor Value Identifier - Code Report</title>
<style>
@page { size: Letter; margin: 0; }
* { box-sizing: border-box; }
body { margin: 0; color: #182230; font: 8.5pt Aptos, Calibri, sans-serif; }
.page { position: relative; width: 8.5in; height: 11in; padding: .48in .58in .58in; overflow: hidden; page-break-after: always; break-after: page; }
.page:last-child { page-break-after: auto; break-after: auto; }
.cover { margin: 0 0 5pt; color: #416689; font-size: 7.5pt; font-weight: bold; letter-spacing: .7pt; }
h1 { margin: 0 0 5pt; color: #1f4e79; font-size: 15pt; }
.summary { margin: 0 0 7pt; line-height: 1.22; }
h2 { margin: 5pt 0 2pt; color: #1f4e79; font-size: 8pt; }
pre { margin: 0 0 4pt; padding: 4pt 6pt; border-left: 2pt solid #7393b3; background: #f2f5f8; white-space: pre-wrap; overflow-wrap: anywhere; }
code { font: 6.8pt/1.08 Consolas, monospace; }
.note { margin-top: 5pt; font-size: 8pt; font-weight: bold; }
footer { position: absolute; right: .58in; bottom: .24in; color: #667085; font-size: 7pt; }
</style></head><body>''' + ''.join(page_html) + '</body></html>'
    HTML_PATH.write_text(html, encoding='utf-8')


def write_pdf() -> None:
    document = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=0.58 * inch,
        leftMargin=0.58 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.55 * inch,
        title='Resistor Value Identifier - Three-Page Code Summary',
        author='Resistor Value Identifier Project',
    )
    title_style = ParagraphStyle(
        'ReportTitle', fontName='Helvetica-Bold', fontSize=14, leading=17,
        textColor=colors.HexColor('#1f4e79'), spaceAfter=5,
    )
    cover_style = ParagraphStyle(
        'CoverLine', fontName='Helvetica-Bold', fontSize=7.5, leading=9,
        textColor=colors.HexColor('#416689'), spaceAfter=4,
    )
    summary_style = ParagraphStyle(
        'Summary', fontName='Helvetica', fontSize=8.2, leading=10,
        textColor=colors.HexColor('#182230'), spaceAfter=5,
    )
    label_style = ParagraphStyle(
        'CodeHeading', fontName='Helvetica-Bold', fontSize=7.8, leading=9,
        textColor=colors.HexColor('#1f4e79'), spaceBefore=3, spaceAfter=2,
    )
    code_style = ParagraphStyle(
        'Code', fontName='Courier', fontSize=6.4, leading=7.4,
        textColor=colors.HexColor('#182230'), leftIndent=5, rightIndent=2,
        borderColor=colors.HexColor('#7393b3'), borderWidth=0.6,
        borderPadding=4, backColor=colors.HexColor('#f2f5f8'), spaceAfter=3,
    )
    note_style = ParagraphStyle(
        'Note', fontName='Helvetica-Bold', fontSize=7.8, leading=9.5,
        textColor=colors.HexColor('#182230'), spaceBefore=4,
    )

    story = []
    for index, page in enumerate(PAGES):
        if index:
            story.append(PageBreak())
        if index == 0:
            story.append(Paragraph('RESISTOR VALUE IDENTIFIER | CONCISE CODE REPORT', cover_style))
        story.append(Paragraph(page['title'], title_style))
        story.append(Paragraph(page['summary'], summary_style))
        for heading, code in page['blocks']:
            story.append(Paragraph(heading.replace('&', '&amp;'), label_style))
            story.append(Preformatted(code.strip(), code_style, maxLineLength=110))
        if page.get('note'):
            story.append(Spacer(1, 3))
            story.append(Paragraph(page['note'], note_style))

    def draw_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.setFillColor(colors.HexColor('#667085'))
        canvas.drawCentredString(letter[0] / 2, 0.25 * inch, f'Resistor Value Identifier | Page {doc.page} of 3')
        canvas.restoreState()

    document.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)


def main() -> None:
    write_docx()
    write_html()
    write_pdf()
    try:
        write_markdown()
    except OSError as error:
        print(f'Warning: unable to refresh {MD_PATH.name}: {error}')
    print(f'Created {PDF_PATH.name}, {DOCX_PATH.name}, and {HTML_PATH.name} with three explicit sections.')


if __name__ == '__main__':
    main()
